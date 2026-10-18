"""Build the downloadable sample BOM offered in the New quote panel.

It reuses the demo BOM's parts (so every result matches the demo catalog) and is laid out to
show every capability: a title block above the header, two non-standard (German) headers for
AI column mapping, data problems, sourcing, every risk type, MOQ/scrap cost effects, alternates
(drop-in, not suitable, cheaper packaging) and a quote with blockers to clear. A Remarks column
says what each interesting line demonstrates, and a Read me sheet walks through the demo.

    uv run python scripts/generate_sample_bom.py
"""

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[3]
SOURCE = ROOT / "data" / "samples" / "motor_controller.xlsx"
TARGET = ROOT / "apps" / "web" / "public" / "samples" / "QuoteLens-sample-BOM.xlsx"

# Header names: "Hersteller" (manufacturer) and "Bestückungsorte" (designators) are not in the
# alias table on purpose, so the AI column mapper has something to do.
HEADERS = [
    "Pos.",
    "Customer P/N",
    "Description",
    "Hersteller",
    "Manufacturer P/N",
    "Qty",
    "Bestückungsorte",
    "UOM",
    "DNP",
    "Target Price (USD)",
    "Remarks",
]

# What each showcase line demonstrates (keyed by designator).
REMARKS = {
    "U1": "Single authorized source. Find alternates: a cheaper reel-packaged drop-in exists.",
    "U2": "Short on authorized stock and a long lead time: see Risks.",
    "U3,U4": "Big price spread between distributors. Its alternate is NOT suitable (gain differs).",
    "U5": "Customer target price is below the market price.",
    "U7": "Not recommended for new designs (end of life approaching).",
    "U11-U13": "OBSOLETE, a blocker. Find alternates, then approve the drop-in to clear it.",
    "U16": "Typo in the part number: QuoteLens suggests TPS54331DDR.",
    "Q1-Q6": "Largest share of the BOM cost.",
    "R16-R21": "6 designators but quantity 5: flagged as a data problem.",
    "R36,R37": "Manufacturer missing: the catalog identifies it.",
    "R44": "Do not populate: excluded from the cost.",
    "C1-C20": "Same part as the next row: the two rows are combined into one line.",
    "C21-C24": "Same part as the previous row (combined).",
    "L1": "MOQ 500 forces excess stock at 100 boards: open Cost impact and Show steps.",
    "J2": "Quantity 'TBD', a blocker. Open the line and correct the quantity.",
    "F1": "No part number, a blocker. Open the line and enter 0451005.MRL.",
}

_STEPS = [
    ("QuoteLens sample BOM", None),
    ("", None),
    ("Use this file to try every QuoteLens feature.", None),
    ("", None),
    (
        "1. Start a new quote",
        "Name it, keep the build quantities 100, 1000, 5000 and pick a target ship date "
        "about 8 weeks from today.",
    ),
    (
        "2. Upload this file",
        "QuoteLens skips the title rows, finds the header and maps the columns; the two "
        "German headers are mapped by AI.",
    ),
    ("3. Confirm the columns", "Check the mapping, then Confirm and analyze."),
    ("4. Overview and Risks", "3 blockers stop the quote. Every risk explains why, with evidence."),
    (
        "5. Fix the blockers",
        "F1: enter a part number. J2: correct the quantity. Hall sensor U11-U13: Find "
        "alternates and approve the drop-in.",
    ),
    (
        "6. Cost impact",
        "Switch build quantities; open Show steps on L1 to see the MOQ excess calculation.",
    ),
    ("7. Ask the copilot", "For example: Which parts threaten delivery by the target date?"),
    ("8. Quote", "Create a quote version, finalize it and download the Excel workbook."),
    ("", None),
    ("The Remarks column in the BOM sheet says what each interesting line demonstrates.", None),
    ("Prices, stock and lifecycle come from the QuoteLens demo catalog (synthetic values).", None),
]
README: list[tuple[str, str | None]] = _STEPS


def build() -> Path:
    source = load_workbook(SOURCE, data_only=True)["BOM"]
    rows = [
        r for r in source.iter_rows(min_row=2, values_only=True) if any(c is not None for c in r)
    ]

    wb = Workbook()
    wb.remove(wb.worksheets[0])
    bom = wb.create_sheet("BOM")
    bom.append(["Northwind Drives: MC48 BLDC motor controller"])
    bom.append(["Revision C, request for quotation"])
    bom.append([])
    bom.append(HEADERS)
    for row in rows:
        item, cpn, desc, mfr, mpn, qty, refdes, uom, dnp, target = row
        bom.append(
            [item, cpn, desc, mfr, mpn, qty, refdes, uom, dnp, target, REMARKS.get(str(refdes))]
        )

    bom["A1"].font = Font(bold=True, size=14)
    header_fill = PatternFill("solid", fgColor="E1EDE7")
    for cell in bom[4]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    for col, width in zip("ABCDEFGHIJK", (6, 13, 34, 20, 22, 6, 16, 6, 6, 12, 70), strict=True):
        bom.column_dimensions[col].width = width
    for r in range(5, bom.max_row + 1):
        if bom[f"K{r}"].value:
            bom[f"K{r}"].fill = PatternFill("solid", fgColor="FBF5E1")
    bom.freeze_panes = "A5"

    readme = wb.create_sheet("Read me")
    for left, right in README:
        readme.append([left, right])
    readme["A1"].font = Font(bold=True, size=14)
    readme.column_dimensions["A"].width = 26
    readme.column_dimensions["B"].width = 110
    for r in range(1, readme.max_row + 1):
        readme[f"B{r}"].alignment = Alignment(wrap_text=True, vertical="top")

    TARGET.parent.mkdir(parents=True, exist_ok=True)
    wb.save(TARGET)
    return TARGET


if __name__ == "__main__":
    path = build()
    print(f"wrote {path.relative_to(ROOT)} ({path.stat().st_size:,} bytes)")
