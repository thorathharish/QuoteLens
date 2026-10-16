"""Generate the demo catalog seed and sample BOMs (Phase 1).

    uv run python scripts/generate_demo_data.py            # writes data/catalog_seed + data/samples
    uv run python scripts/generate_demo_data.py --check    # fail if committed files are stale

Output is deterministic: synthetic values derive from SHA-256 of the part number, so
re-running produces identical catalog and CSV files. Every row is validated against the
catalog row models before it is written.
"""

import argparse
import csv
import hashlib
import io
import json
import sys
from collections.abc import Iterable
from datetime import UTC, date, datetime
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from demo_data.boms import (
    IOT_SENSOR,
    MESSY_FORMAT,
    MESSY_HEADERS,
    MESSY_TITLE_BLOCK,
    MOTOR_CONTROLLER,
    BomRow,
)
from demo_data.parts import CROSSREFS, PARTS, PartSpec
from openpyxl import Workbook
from openpyxl.styles import Font
from pydantic import BaseModel

from app.catalog.rows import CatalogComponent, CatalogCrossRef, CatalogOffer, CatalogPriceBreak
from app.domain.enums import SourceMode
from app.domain.mpn import normalize_mpn

REPO_ROOT = Path(__file__).resolve().parents[3]
CATALOG_DIR = REPO_ROOT / "data" / "catalog_seed"
SAMPLES_DIR = REPO_ROOT / "data" / "samples"

SNAPSHOT_DATE = date(2026, 10, 17)
FETCHED_AT = datetime(2026, 10, 17, 6, 0, tzinfo=UTC)
SOURCE = "QuoteLens demo catalog (synthetic commercial data)"
CURRENCY = "USD"
DISTRIBUTORS = (("DigiKey", "DK"), ("Mouser", "MOU"), ("Arrow", "ARW"))
PASSIVE = {"resistor", "capacitor", "ferrite", "thermistor", "led", "diode"}
LONG_LEAD_DAYS = 112  # 16 weeks
BLOCKER_RULES = {"D01", "D04", "L01"}  # seeded rules whose default severity is BLOCKER
BUILD_QUANTITIES = [100, 1000, 5000]
TARGET_SHIP_OFFSET_DAYS = 56


# ── deterministic synthetic values ───────────────────────────


def _unit(*parts: str) -> float:
    """Stable pseudo-random number in [0, 1) derived from the inputs."""
    digest = hashlib.sha256("|".join(parts).encode()).digest()
    return int.from_bytes(digest[:8], "big") / 2**64


def _money(value: Decimal) -> Decimal:
    return max(value, Decimal("0.0005")).quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP)


def _breaks(
    spec: PartSpec, factor: Decimal, schedule: list[tuple[int, str]]
) -> list[CatalogPriceBreak]:
    breaks: list[CatalogPriceBreak] = []
    previous: Decimal | None = None
    for qty, multiplier in schedule:
        price = _money(spec.unit_price * Decimal(multiplier) * factor)
        if previous is not None:
            price = min(price, previous)
        breaks.append(CatalogPriceBreak(qty=qty, price=price))
        previous = price
    return breaks


def _reel_qty(spec: PartSpec) -> int:
    if spec.reel_qty:
        return spec.reel_qty
    if spec.category in PASSIVE:
        return 10_000 if spec.package in {"0402", "0603"} else 4_000
    return 3_000 if spec.package and spec.package.startswith("SOT") else 2_500


def _cut_tape_schedule(spec: PartSpec) -> list[tuple[int, str]]:
    if spec.category in PASSIVE:
        return [(1, "1"), (10, "0.62"), (100, "0.24"), (1000, "0.085")]
    return [(1, "1"), (10, "0.9"), (25, "0.86"), (100, "0.76"), (250, "0.71"), (1000, "0.63")]


def _packaging(spec: PartSpec) -> str:
    if spec.category in {"connector", "switch", "fuse"} and spec.package == "THT":
        return "Bulk"
    if spec.category == "module":
        return "Tray"
    return "Cut Tape"


def _sku(prefix: str, spec: PartSpec, suffix: str) -> str:
    key = "".join(ch if ch.isalnum() else "-" for ch in spec.mpn.upper())
    return f"DEMO-{prefix}-{key}-{suffix}"


def _offer(
    spec: PartSpec,
    distributor: str,
    sku: str,
    *,
    stock: int,
    moq: int,
    mult: int,
    lead_days: int | None,
    packaging: str,
    breaks: list[CatalogPriceBreak],
    authorized: bool = True,
) -> CatalogOffer:
    return CatalogOffer(
        mpn_norm=_norm(spec.mpn),
        distributor=distributor,
        sku=sku,
        manufacturer=spec.manufacturer,
        stock=stock,
        moq=moq,
        mult=mult,
        lead_days=lead_days,
        packaging=packaging,
        price_breaks=breaks,
        currency=CURRENCY,
        authorized=authorized,
        lifecycle=spec.lifecycle,
        fetched_at=FETCHED_AT,
        mode=SourceMode.DEMO_SNAPSHOT,
    )


def offers_for(spec: PartSpec) -> list[CatalogOffer]:
    offers: list[CatalogOffer] = []
    reel = _reel_qty(spec)
    for distributor, prefix in DISTRIBUTORS:
        h = _unit(spec.mpn, distributor)
        factor = Decimal(1) + (Decimal(str(round(h, 4))) - Decimal("0.5")) * Decimal("0.16")
        if spec.profile == "price_outlier" and distributor == "Mouser":
            factor = Decimal("2.3")
        lead: int | None = 42 + int(_unit(spec.mpn, distributor, "lead") * 57)
        stock = 20_000 + int(h * 180_000) if spec.category in PASSIVE else 800 + int(h * 12_000)

        if spec.profile == "single_source" and distributor != "DigiKey":
            continue
        if spec.profile == "long_lead":
            stock = {"DigiKey": 37, "Mouser": 0, "Arrow": 12}[distributor]
            lead = LONG_LEAD_DAYS
        if spec.profile == "no_stock":
            stock, lead = 0, None
        if spec.profile == "single_source":
            stock = 2_400

        if spec.profile == "reel_only":
            schedule = [(reel, "0.7"), (reel * 2, "0.64"), (reel * 5, "0.58")]
            offers.append(
                _offer(
                    spec,
                    distributor,
                    _sku(prefix, spec, "TR"),
                    stock=stock,
                    moq=reel,
                    mult=reel,
                    lead_days=lead,
                    packaging="Tape & Reel",
                    breaks=_breaks(spec, factor, schedule),
                )
            )
            continue

        offers.append(
            _offer(
                spec,
                distributor,
                _sku(prefix, spec, "CT"),
                stock=stock,
                moq=1,
                mult=1,
                lead_days=lead,
                packaging=_packaging(spec),
                breaks=_breaks(spec, factor, _cut_tape_schedule(spec)),
            )
        )
        if distributor == "DigiKey" and spec.category in PASSIVE and spec.profile == "standard":
            reel_schedule = [(reel, "0.06")]
            offers.append(
                _offer(
                    spec,
                    distributor,
                    _sku(prefix, spec, "TR"),
                    stock=stock * 2,
                    moq=reel,
                    mult=reel,
                    lead_days=lead,
                    packaging="Tape & Reel",
                    breaks=_breaks(spec, factor, reel_schedule),
                )
            )

    if spec.profile == "no_stock":
        # Unauthorized open-market stock: available, but a counterfeit/traceability risk.
        offers.append(
            _offer(
                spec,
                "Open-market broker (demo)",
                _sku("BRK", spec, "LOT"),
                stock=1_500,
                moq=100,
                mult=1,
                lead_days=5,
                packaging="Loose",
                breaks=_breaks(spec, Decimal("2.6"), [(100, "1")]),
                authorized=False,
            )
        )
    return offers


def _norm(mpn: str) -> str:
    key = normalize_mpn(mpn)
    if key is None:
        raise ValueError(f"Part with empty MPN: {mpn!r}")
    return key


def build_catalog() -> tuple[list[CatalogComponent], list[CatalogOffer], list[CatalogCrossRef]]:
    by_mpn = {spec.mpn: spec for spec in PARTS}
    if len(by_mpn) != len(PARTS):
        raise ValueError("Duplicate MPN in PARTS")
    components = [
        CatalogComponent(
            mpn_norm=_norm(spec.mpn),
            mpn=spec.mpn,
            manufacturer=spec.manufacturer,
            category=spec.category,
            description=spec.description,
            package=spec.package,
            params=dict(spec.params),
            lifecycle=spec.lifecycle,
            ltb_date=spec.ltb_date,
            rohs=True,
            source=SOURCE,
            snapshot_date=SNAPSHOT_DATE,
        )
        for spec in PARTS
    ]
    offers = [offer for spec in PARTS for offer in offers_for(spec)]
    crossrefs = []
    for original, alternate, relation, note in CROSSREFS:
        alt = by_mpn[alternate]  # KeyError here means an alternate is missing from PARTS
        crossrefs.append(
            CatalogCrossRef(
                mpn_norm=_norm(by_mpn[original].mpn),
                alt_mpn_norm=_norm(alt.mpn),
                alt_mpn=alt.mpn,
                alt_manufacturer=alt.manufacturer,
                relation=relation,
                source="QuoteLens demo catalog (curated)",
                note=note,
            )
        )
    return components, offers, crossrefs


# ── writers ──────────────────────────────────────────────────


def _jsonl(rows: Iterable[BaseModel]) -> str:
    return "".join(row.model_dump_json(exclude_none=True) + "\n" for row in rows)


def _bom_csv(rows: list[BomRow]) -> str:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow(["Reference", "Quantity", "Manufacturer", "MPN", "Description"])
    for row in rows:
        writer.writerow([row.refdes, row.qty, row.manufacturer, row.mpn, row.description])
    return buffer.getvalue()


def _workbook() -> Workbook:
    wb = Workbook()
    wb.properties.creator = "QuoteLens demo data generator"
    wb.properties.created = datetime(2026, 10, 17)
    return wb


def _motor_xlsx(path: Path) -> None:
    wb = _workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "BOM"
    headers = [
        "Item",
        "Customer P/N",
        "Description",
        "Manufacturer",
        "Manufacturer P/N",
        "Qty",
        "Ref Des",
        "UOM",
        "DNP",
        "Target Price (USD)",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for item, row in enumerate(MOTOR_CONTROLLER, start=1):
        qty: int | str = int(row.qty) if row.qty.isdigit() else row.qty
        price: float | str = float(row.target_price) if row.target_price else ""
        ws.append(
            [
                item,
                f"MC48-{item:03d}",
                row.description,
                row.manufacturer,
                row.mpn,
                qty,
                row.refdes,
                row.uom,
                row.dnp,
                price,
            ]
        )
    notes = wb.create_sheet("Notes")
    notes.append(["Assembly", "MC48 BLDC motor controller, Rev B"])
    notes.append(["Build quantities", ", ".join(str(q) for q in BUILD_QUANTITIES)])
    notes.append(["Data", "Sample BOM for the QuoteLens demo"])
    wb.save(path)


def _messy_xlsx(path: Path) -> None:
    wb = _workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Sheet1"
    for line in MESSY_TITLE_BLOCK:
        ws.append([line])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(MESSY_HEADERS))
    ws.append([])
    ws.append(MESSY_HEADERS)
    for item, row in enumerate(MESSY_FORMAT, start=1):
        fitted = "N" if row.dnp == "N" else "Y"
        ws.append(
            [
                item,
                row.cpn,
                row.description,
                row.manufacturer,
                row.mpn,
                int(row.qty),
                row.uom,
                row.refdes,
                fitted,
                row.extra.get("Remarks", ""),
            ]
        )
    wb.save(path)


def _expected(
    file: str,
    rows: list[BomRow],
    *,
    sheet: str | None,
    header_row: int,
    consolidated_lines: int,
    open_blockers: int,
) -> str:
    seeded = [
        {
            "row": index,
            "refdes": row.refdes,
            "mpn": row.mpn or None,
            "match": row.expect.match or None,
            "issues": list(row.expect.issues),
            "risks": list(row.expect.risks),
            "note": row.expect.note,
        }
        for index, row in enumerate(rows, start=1)
        if row.expect is not None
    ]
    manifest = {
        "file": file,
        "sheet": sheet,
        "header_row": header_row,
        "data_rows": len(rows),
        "lines_after_consolidation": consolidated_lines,
        "build_quantities": BUILD_QUANTITIES,
        "target_ship_offset_days": TARGET_SHIP_OFFSET_DAYS,
        "open_blockers": open_blockers,
        "seeded_rows": seeded,
        "unseeded_rows_expectation": "MATCHED with no data issues and no seeded risks",
        "totals": None,
        "_totals_note": "Per-volume totals are frozen in Phase 4, once the impact engine exists.",
    }
    return json.dumps(manifest, indent=2, ensure_ascii=False) + "\n"


def render_text_outputs() -> dict[Path, str]:
    components, offers, crossrefs = build_catalog()
    blockers = sum(
        1
        for row in MOTOR_CONTROLLER
        if row.expect is not None and BLOCKER_RULES.intersection(row.expect.risks)
    )
    return {
        CATALOG_DIR / "components.jsonl": _jsonl(components),
        CATALOG_DIR / "offers.jsonl": _jsonl(offers),
        CATALOG_DIR / "crossrefs.jsonl": _jsonl(crossrefs),
        SAMPLES_DIR / "iot_sensor.csv": _bom_csv(IOT_SENSOR),
        SAMPLES_DIR / "motor_controller.expected.json": _expected(
            "motor_controller.xlsx",
            MOTOR_CONTROLLER,
            sheet="BOM",
            header_row=1,
            consolidated_lines=len(MOTOR_CONTROLLER) - 1,
            open_blockers=blockers,
        ),
        SAMPLES_DIR / "iot_sensor.expected.json": _expected(
            "iot_sensor.csv",
            IOT_SENSOR,
            sheet=None,
            header_row=1,
            consolidated_lines=len(IOT_SENSOR),
            open_blockers=0,
        ),
        SAMPLES_DIR / "messy_format.expected.json": _expected(
            "messy_format.xlsx",
            MESSY_FORMAT,
            sheet="Sheet1",
            header_row=len(MESSY_TITLE_BLOCK) + 2,
            consolidated_lines=len(MESSY_FORMAT),
            open_blockers=0,
        ),
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--check", action="store_true", help="fail if committed outputs are stale")
    args = parser.parse_args()

    outputs = render_text_outputs()
    if args.check:
        stale = [
            p
            for p, text in outputs.items()
            if not p.exists() or p.read_text(encoding="utf-8") != text
        ]
        for path in stale:
            print(f"stale: {path.relative_to(REPO_ROOT)}")
        return 1 if stale else 0

    CATALOG_DIR.mkdir(parents=True, exist_ok=True)
    SAMPLES_DIR.mkdir(parents=True, exist_ok=True)
    for path, text in outputs.items():
        path.write_text(text, encoding="utf-8", newline="\n")
        print(f"wrote {path.relative_to(REPO_ROOT)}")
    for name, writer in (
        ("motor_controller.xlsx", _motor_xlsx),
        ("messy_format.xlsx", _messy_xlsx),
    ):
        writer(SAMPLES_DIR / name)
        print(f"wrote data/samples/{name}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
