"""Phase 7 end to end: the full demo story from seed to finalized, exported quote."""

import io
from typing import Any

from httpx import AsyncClient
from openpyxl import load_workbook

from app.container import Container
from app.infra.jobbus import InlineJobBus


async def _drain(container: Container) -> None:
    assert isinstance(container.job_bus, InlineJobBus)
    await container.job_bus.drain()


async def _lines(client: AsyncClient, pid: str) -> dict[str, dict[str, Any]]:
    items = (await client.get(f"/api/v1/projects/{pid}/lines", params={"limit": 500})).json()[
        "items"
    ]
    return {(line["mpn"] or line["refdes"][0]): line for line in items}


async def test_demo_story_from_seed_to_finalized_quote(
    api_client: AsyncClient, container: Container
) -> None:
    seeded = await api_client.post("/api/v1/demo/projects", json={"scenario": "motor_controller"})
    assert seeded.status_code == 201, seeded.text
    pid = seeded.json()["project_id"]
    await _drain(container)
    project = (await api_client.get(f"/api/v1/projects/{pid}")).json()
    assert project["data_mode"] == "demo_snapshot"
    assert project["status"] == "IN_REVIEW"

    draft = (
        await api_client.post(
            f"/api/v1/projects/{pid}/quotes",
            json={"note": "First pass"},
            headers={"X-Demo-Role": "QUOTE_ENGINEER"},
        )
    ).json()
    assert draft["version"] == 1
    assert draft["status"] == "DRAFT"
    assert draft["readiness"]["ready"] is False
    assert len(draft["readiness"]["open_blockers"]) == 3
    assert [v["volume"] for v in draft["volumes"]] == [100, 1000, 5000]
    assert len(draft["lines"]) == 59

    blocked = await api_client.post(
        f"/api/v1/projects/{pid}/quotes/1/finalize", json={"note": "Send it"}
    )
    assert blocked.status_code == 409
    assert blocked.json()["error"]["code"] == "QUOTE_HAS_OPEN_BLOCKERS"
    assert len(blocked.json()["error"]["details"]["open_blockers"]) == 3

    lines = await _lines(api_client, pid)
    # Blocker 1: missing MPN → corrected via a general line edit
    fixed = await api_client.patch(
        f"/api/v1/projects/{pid}/lines/{lines['F1']['id']}",
        json={
            "changes": {"mpn": "0451005.MRL", "manufacturer": "Littelfuse"},
            "reason": "Customer confirmed",
        },
    )
    assert fixed.status_code == 200, fixed.text
    assert fixed.json()["line"]["match"]["status"] == "MATCHED"
    # Blocker 2: quantity "TBD" → user-entered quantity with its own evidence
    qty = await api_client.patch(
        f"/api/v1/projects/{pid}/lines/{lines['0430450200']['id']}",
        json={"changes": {"qty_per": "1"}, "reason": "One power connector per board"},
        headers={"X-Demo-Role": "QUOTE_ENGINEER"},
    )
    qty_fact = qty.json()["line"]["qty_per"]
    assert qty_fact["provenance"] == "ACTUAL"
    (ev_id,) = qty_fact["evidence_ids"]
    evidence = (await api_client.get(f"/api/v1/projects/{pid}/evidence/{ev_id}")).json()
    assert evidence["kind"] == "USER_OVERRIDE"
    assert evidence["payload_excerpt"]["old"] == "None"
    assert evidence["payload_excerpt"]["reason"] == "One power connector per board"
    # Blocker 3: obsolete hall sensor → approve the drop-in alternate
    hall = lines["A1101ELHLT-T"]["id"]
    await api_client.post(f"/api/v1/projects/{pid}/lines/{hall}/alternatives/search")
    await _drain(container)
    (alt,) = (
        await api_client.get(f"/api/v1/projects/{pid}/alternatives", params={"line_id": hall})
    ).json()["items"]
    await api_client.post(
        f"/api/v1/projects/{pid}/alternatives/{alt['id']}/decision",
        json={"decision": "APPROVE", "note": "Drop-in successor"},
    )

    assert (await api_client.get(f"/api/v1/projects/{pid}")).json()["status"] == "QUOTE_READY"
    final = await api_client.post(
        f"/api/v1/projects/{pid}/quotes/1/finalize",
        json={"note": "Approved for customer"},
        headers={"X-Demo-Role": "MANAGER"},
    )
    assert final.status_code == 200, final.text
    assert final.json()["status"] == "FINALIZED"
    assert final.json()["finalized_by"] == "MANAGER"
    assert final.json()["readiness"]["ready"] is True
    assert (await api_client.get(f"/api/v1/projects/{pid}")).json()["status"] == "FINALIZED"

    again = await api_client.post(
        f"/api/v1/projects/{pid}/quotes/1/finalize", json={"note": "again"}
    )
    assert again.status_code == 409
    locked = await api_client.patch(
        f"/api/v1/projects/{pid}/lines/{lines['F1']['id']}",
        json={"changes": {"description": "x"}, "reason": "late edit"},
    )
    assert locked.status_code == 409  # finalized projects are read-only until reopened

    export = await api_client.get(f"/api/v1/projects/{pid}/quotes/1/export")
    assert export.status_code == 200
    assert "attachment" in export.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(export.content))
    assert wb.sheetnames == ["Summary", "Lines", "Risks", "Assumptions"]
    summary_text = " ".join(
        str(c.value) for row in wb["Summary"].iter_rows() for c in row if c.value
    )
    assert "DEMO SNAPSHOT" in summary_text
    assert "Ready to quote" in summary_text
    assert wb["Lines"].max_row == 60  # header + 59 lines
    assert wb["Lines"]["L1"].value == "Price provenance"

    reopened = await api_client.post(
        f"/api/v1/projects/{pid}/reopen", json={"reason": "Customer changed volumes"}
    )
    assert reopened.status_code == 204
    v2 = (await api_client.post(f"/api/v1/projects/{pid}/quotes")).json()
    assert v2["version"] == 2
    old = await api_client.post(
        f"/api/v1/projects/{pid}/quotes/1/finalize", json={"note": "old one"}
    )
    assert old.status_code == 409

    audit = (await api_client.get(f"/api/v1/projects/{pid}/audit")).json()["items"]
    actions = [e["action"] for e in audit]
    assert {
        "QUOTE_CREATED",
        "QUOTE_FINALIZED",
        "LINE_EDITED",
        "ALTERNATE_APPROVED",
        "PROJECT_REOPENED",
    } <= set(actions)
    assert audit[0]["at"] >= audit[-1]["at"]  # newest first


async def test_validation_errors(api_client: AsyncClient, container: Container) -> None:
    bad = await api_client.post("/api/v1/demo/projects", json={"scenario": "nope"})
    assert bad.status_code == 422

    pid = (await api_client.post("/api/v1/demo/projects", json={"scenario": "iot_sensor"})).json()[
        "project_id"
    ]
    await _drain(container)
    line = next(iter((await _lines(api_client, pid)).values()))
    same = await api_client.patch(
        f"/api/v1/projects/{pid}/lines/{line['id']}",
        json={"changes": {"mpn": line["mpn"]}, "reason": "no-op"},
    )
    assert same.status_code == 422
    quote = (await api_client.post(f"/api/v1/projects/{pid}/quotes")).json()
    assert quote["readiness"]["ready"] is True  # the clean BOM has no blockers
    empty = (
        await api_client.post("/api/v1/projects", json={"name": "E", "build_quantities": [1]})
    ).json()
    assert (await api_client.post(f"/api/v1/projects/{empty['id']}/quotes")).status_code == 422
    assert (
        await api_client.post(f"/api/v1/projects/{empty['id']}/reopen", json={"reason": "why not"})
    ).status_code == 409
