"""Guards the Phase 1 demo data: schema parity, referential integrity, seeded scenarios."""

import csv
import json
import subprocess
import sys
from decimal import Decimal
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook
from pydantic import BaseModel

from app.catalog.rows import CatalogComponent, CatalogCrossRef, CatalogOffer
from app.domain.enums import Lifecycle, SourceMode
from app.domain.mpn import normalize_mpn

API_DIR = Path(__file__).resolve().parents[2]
REPO_ROOT = API_DIR.parents[1]
SEED = REPO_ROOT / "data" / "catalog_seed"
SAMPLES = REPO_ROOT / "data" / "samples"
SCHEMAS = REPO_ROOT / "infra" / "bq" / "schemas"


def _load[M: BaseModel](name: str, model: type[M]) -> list[M]:
    lines = (SEED / name).read_text(encoding="utf-8").splitlines()
    return [model.model_validate_json(line) for line in lines if line.strip()]


@pytest.fixture(scope="module")
def components() -> dict[str, CatalogComponent]:
    return {c.mpn_norm: c for c in _load("components.jsonl", CatalogComponent)}


@pytest.fixture(scope="module")
def offers() -> list[CatalogOffer]:
    return _load("offers.jsonl", CatalogOffer)


def _offers_for(offers: list[CatalogOffer], mpn: str) -> list[CatalogOffer]:
    return [o for o in offers if o.mpn_norm == mpn]


def _manifest(name: str) -> dict[str, Any]:
    data: dict[str, Any] = json.loads((SAMPLES / name).read_text(encoding="utf-8"))
    return data


def test_committed_seed_is_up_to_date() -> None:
    result = subprocess.run(
        [sys.executable, "scripts/generate_demo_data.py", "--check"],
        cwd=API_DIR,
        capture_output=True,
        text=True,
        check=False,
    )
    assert result.returncode == 0, f"Regenerate demo data:\n{result.stdout}{result.stderr}"


@pytest.mark.parametrize(
    ("table", "model"),
    [
        ("components", CatalogComponent),
        ("offers_cache", CatalogOffer),
        ("crossrefs", CatalogCrossRef),
    ],
)
def test_bigquery_schema_matches_row_model(table: str, model: type[BaseModel]) -> None:
    schema = json.loads((SCHEMAS / f"{table}.json").read_text(encoding="utf-8"))
    assert {field["name"] for field in schema} == set(model.model_fields)


def test_every_offer_and_crossref_points_at_a_component(
    components: dict[str, CatalogComponent], offers: list[CatalogOffer]
) -> None:
    assert {o.mpn_norm for o in offers} <= set(components)
    for ref in _load("crossrefs.jsonl", CatalogCrossRef):
        assert ref.mpn_norm in components
        assert ref.alt_mpn_norm in components
    assert all(o.mode is SourceMode.DEMO_SNAPSHOT for o in offers)
    assert all(o.sku.startswith("DEMO-") for o in offers), "synthetic SKUs must be marked"


def _bom_mpns() -> dict[str, list[str]]:
    motor = load_workbook(SAMPLES / "motor_controller.xlsx", read_only=True)["BOM"]
    messy = load_workbook(SAMPLES / "messy_format.xlsx", read_only=True)["Sheet1"]
    with (SAMPLES / "iot_sensor.csv").open(encoding="utf-8", newline="") as handle:
        iot = [row["MPN"] for row in csv.DictReader(handle)]
    return {
        "motor_controller.expected.json": [
            str(r[4] or "") for r in motor.iter_rows(min_row=2, values_only=True)
        ],
        "messy_format.expected.json": [
            str(r[4] or "") for r in messy.iter_rows(min_row=6, values_only=True)
        ],
        "iot_sensor.expected.json": iot,
    }


@pytest.mark.parametrize("manifest_name", list(_bom_mpns()))
def test_bom_parts_resolve_unless_seeded_otherwise(
    manifest_name: str, components: dict[str, CatalogComponent]
) -> None:
    manifest = _manifest(manifest_name)
    mpns = _bom_mpns()[manifest_name]
    assert len(mpns) == manifest["data_rows"]
    seeded = {row["row"]: row for row in manifest["seeded_rows"]}
    for row_no, mpn in enumerate(mpns, start=1):
        expected_match = seeded.get(row_no, {}).get("match", "MATCHED")
        found = normalize_mpn(mpn) in components
        if expected_match == "MATCHED":
            assert found, f"{manifest_name} row {row_no}: {mpn} missing from catalog"
        else:
            assert not found, f"{manifest_name} row {row_no}: {mpn} should not resolve"


def test_motor_controller_seeds_every_scenario(
    components: dict[str, CatalogComponent], offers: list[CatalogOffer]
) -> None:
    manifest = _manifest("motor_controller.expected.json")
    seeded_risks = {risk for row in manifest["seeded_rows"] for risk in row["risks"]}
    assert seeded_risks >= {
        "D01",
        "D02",
        "D03",
        "D04",
        "D05",
        "L01",
        "L02",
        "S01",
        "S02",
        "S03",
        "P01",
        "P02",
        "P03",
    }
    assert manifest["open_blockers"] == 3

    # L01: obsolete with no authorized stock; only an unauthorized broker has parts.
    hall = _offers_for(offers, "A1101ELHLT-T")
    assert components["A1101ELHLT-T"].lifecycle is Lifecycle.OBSOLETE
    assert all(o.stock == 0 for o in hall if o.authorized)
    assert any(not o.authorized and (o.stock or 0) > 0 for o in hall)
    # L02
    assert components["MCP2551-I/SN"].lifecycle is Lifecycle.NRND
    # S02: single authorized source
    assert len(_offers_for(offers, "STM32G431CBT6")) == 1
    # S01/S03: 16-week lead and less stock than one build needs
    driver = _offers_for(offers, "DRV8323RSRGZR")
    assert all(o.lead_days == 112 for o in driver)
    assert sum(o.stock or 0 for o in driver) < 100
    # P01: one distributor is far above the median at qty 1
    prices = sorted(o.price_breaks[0].price for o in _offers_for(offers, "INA240A1PWR"))
    assert prices[-1] > prices[len(prices) // 2] * Decimal("1.4")
    # P03: reel-only MOQ
    assert all(o.moq == 500 and o.mult == 500 for o in _offers_for(offers, "SRP1265A-220M"))
    # D03 near-miss: the intended part exists, the typo does not
    assert "TPS54331DDR" in components
    assert "TPS5431DDR" not in components


def test_messy_format_header_row_is_below_title_block() -> None:
    manifest = _manifest("messy_format.expected.json")
    sheet = load_workbook(SAMPLES / "messy_format.xlsx", read_only=True)["Sheet1"]
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[manifest["header_row"] - 1][4] == "Maker Part No."


def test_mpn_normalization_keeps_punctuation() -> None:
    assert normalize_mpn("  bat54s,215 ") == "BAT54S,215"
    assert normalize_mpn("mcp2561-e/sn​") == "MCP2561-E/SN"
    assert normalize_mpn("   ") is None
    assert normalize_mpn(None) is None
